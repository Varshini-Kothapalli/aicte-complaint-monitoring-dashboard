"""Package cleaned ACSS + synthetic data into one workbook and a CSV bundle.
Usage: python3 -I package.py <work_dir> <out_dir>
"""
import sys, os, zipfile, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

W, OUT = sys.argv[1], sys.argv[2]
os.makedirs(f"{OUT}/csv", exist_ok=True)
acss = pd.read_pickle(f"{W}/acss_clean.pkl")
stats = pd.read_csv(f"{W}/acss_stats.csv", index_col=0)["value"]
syn = {n: pd.read_pickle(f"{W}/syn_{n}.pkl") for n in ["institutions", "institution_alias_tests", "bureaus", "users", "category_matrix",
                                                       "cases", "case_actions", "transfers", "overlap_ground_truth"]}
def CL(df, name):
    return get_column_letter(list(df.columns).index(name) + 1)


A_ = lambda n: CL(acss, n)
C_ = lambda n: CL(syn["cases"], n)
G_ = lambda n: CL(syn["overlap_ground_truth"], n)
griev = acss[acss.record_class == "grievance_against_institution"].copy()

wb = Workbook()
normal = wb._named_styles["Normal"]
normal.font = Font(name="Arial", size=10)
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3A5F")
H1 = Font(name="Arial", size=14, bold=True, color="1F3A5F")
H2 = Font(name="Arial", size=11, bold=True, color="1F3A5F")
B = Font(name="Arial", size=10, bold=True)
NOTE = Font(name="Arial", size=9, italic=True, color="555555")
WRAP = Alignment(wrap_text=True, vertical="top")
SYN_FILL = PatternFill("solid", fgColor="F3E8FF")
IMP_FILL = PatternFill("solid", fgColor="FEF3C7")


def text_sheet(ws, rows, widths=(28, 110)):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
        style = row[0] if row else ""
    return ws


def data_sheet(name, df, fill=None, widths=None, wide_cols=()):
    ws = wb.create_sheet(name)
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.tz_localize(None) if getattr(df[col].dt, "tz", None) else df[col]
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NaT else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value[:1] in "=+-@" and cell.data_type == "f":
                cell.data_type = "s"          # literal text such as "=+ button disabled", not a formula
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, HFILL, Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    for i, col in enumerate(df.columns, 1):
        L = get_column_letter(i)
        if col in wide_cols:
            ws.column_dimensions[L].width = 70
        else:
            sample = df[col].astype(str).head(200).str.len()
            ws.column_dimensions[L].width = max(10, min(32, int(max(len(col), sample.quantile(0.9) if len(sample) else 10)) + 2))
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            fmt = "yyyy-mm-dd hh:mm" if col in ("created_on", "ts", "last_action_at", "requested_at", "accepted_at") else "yyyy-mm-dd"
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ref = f"A1:{get_column_letter(df.shape[1])}{df.shape[0] + 1}"
    t = Table(displayName=name.replace(" ", "_"), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    if fill is not None:
        ws.sheet_properties.tabColor = fill
    return ws


# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
cases = syn["cases"]
rows = [
    ("SIH 2026 — AICTE Institution Complaint Monitoring: prototype dataset", ""),
    ("", ""),
    ("Verdict on the original dataset", "Usable as supporting data, NOT as the core case data. The ACSS file is a genuine public dataset from AICTE's Centralised Support System (Jan–Sep 2022), but almost all of it is institutions asking AICTE for help with the approval portal (EoA download, DSC signing, fee payment, faculty records). It has no institution ID, Bureau, officer, status or deadline, and only about 10 distinct complaints are actually against an institution."),
    ("What was changed", "1) Removed 3,174 exact duplicate rows and 3 test rows (14,272 → 11,095, matching the ~11,115 records the publisher states).  2) Parsed both date formats; flagged rows where Excel may have swapped day and month.  3) Redacted personal and identifying details still present in the text (names, institution names, emails, phones, bank/IFSC/PAN, IDs, addresses, URLs).  4) Classified every row by what it really is, and reviewed every row flagged as a grievance against an institution.  5) Labelled every row IMPORTED_UNVERIFIED with source and DOI."),
    ("What was added", "A clearly-SYNTHETIC monitoring layer the real data cannot provide: 61 fictional institutions (letter-triplet names, SYN- IDs), identity-resolution test names, 3 Bureaus, 27 role-numbered users, 809 cases with full action histories, transfers and escalations, and 214 labelled case pairs for measuring the overlap engine."),
    ("", ""),
    ("Sheets", ""),
    ("Summary", "Live counts (formulas) over the data sheets."),
    ("Data_Dictionary", "Every column of every sheet."),
    ("ACSS_Clean", "11,095 cleaned, redacted, classified ACSS rows. Provenance IMPORTED_UNVERIFIED. Use for analytics and as realistic text."),
    ("ACSS_Grievances", "The 28 rows (10 distinct matters) that are genuine grievances against an institution, after review. Includes two rows describing a complaint forwarded from the complaint cell to the Approval Bureau without follow-up — real-world evidence for the problem statement."),
    ("Institutions", "SYNTHETIC master list. Two institutions deliberately share the name 'PQR Institute of Technology' (Pune and Belagavi)."),
    ("Identity_Tests", "Raw institution names with the expected identity decision (VERIFIED / VERIFICATION_QUEUE / UNRESOLVED). Use to test identity resolution."),
    ("Bureaus, Users", "SYNTHETIC. SLA days are demo configuration, not AICTE timelines. No real officer names."),
    ("Category_Matrix", "Related-category pairs used by the rule check (e.g. fee_refund ↔ fee_regulation_compliance)."),
    ("Cases", "SYNTHETIC cases with every mandatory field: case ID, receipt date, category, status, Bureau, officer, last action, next action, expected disposal date. Demo 'today' = 2026-10-31."),
    ("Case_Actions, Transfers", "Chronological history per case: creation, assignment, status changes, documents, transfers, escalations."),
    ("Overlap_Ground_Truth", "214 labelled case pairs: same_matter (101, 40 of them sharing a reference number), different_matter (98, 30 of them hard negatives in related categories), duplicate_submission (15). expected_min_band is the minimum band a correct engine should reach."),
    ("", ""),
    ("Demo script", "ABC Engineering College (SYN-INST-0001): SYN-PGRC-2026-0001 (PGRC, fee refund, Under Process) and SYN-RB-2026-0002 (RB, refund-norm compliance, Action Initiated) are the same matter in different words → should raise 'Potential Overlapping Cases — Coordination Required'. SYN-AB-2026-0003 (AB, lab deficiency) at the same institution must NOT be flagged."),
    ("", ""),
    ("Source", "ACSS (AICTE Centralized Support System) Data Set — Manoj Singh, Subhash Panwar, Sanju Choudhary; Harvard Dataverse, 2024-01-05. https://doi.org/10.7910/DVN/EITGBB  (licence not stated on the index page: check the Dataverse page before redistribution and cite it in your deck)."),
    ("Limitations", "Redaction is automatic (rules + spaCy NER) and was spot-checked, not proof-read row by row: a few names may remain and some ordinary words were over-redacted. Classification outside the grievance class is rule-based and unreviewed (column review_status). About 4,500 dates with day ≤ 12 may have day and month swapped (date_confidence = medium); monthly totals are barely affected."),
    ("Never present as real", "Anything with provenance SYNTHETIC. Synthetic IDs start with SYN- and synthetic e-mails end in @demo.invalid on purpose."),
]
text_sheet(ws, rows)
ws["A1"].font = H1
for r in range(3, len(rows) + 1):
    if ws.cell(r, 1).value:
        ws.cell(r, 1).font = B
ws.sheet_properties.tabColor = "1F3A5F"

# ------------------------------------------------------------------ Summary (formulas)
sm = wb.create_sheet("Summary")
sm.column_dimensions["A"].width = 46
sm.column_dimensions["B"].width = 14
sm.column_dimensions["C"].width = 60
sm["A1"] = "Dataset summary (live formulas over the data sheets)"; sm["A1"].font = H1
r = 3


def head(t):
    global r
    sm.cell(r, 1, t).font = H2; r += 1


def line(label, formula, note=""):
    global r
    sm.cell(r, 1, label); sm.cell(r, 2, formula); sm.cell(r, 3, note).font = NOTE; r += 1


head("Original ACSS file → cleaned")
for k, lbl, note in [("rows_in_file", "Rows in original file", "Value recorded at import (original file not embedded)"),
                     ("exact_duplicates_removed", "Exact duplicate rows removed", "Same module, nature, subject, description, type, timestamp"),
                     ("junk_or_test_removed", "Test / junk rows removed", "e.g. description 'test'")]:
    line(lbl, int(stats[k]), note)
line("Rows kept (ACSS_Clean)", "=COUNTA(ACSS_Clean!A:A)-1")
KEPT_ROW = r - 1
line("Resubmissions of the same complaint (kept, flagged)", f'=COUNTIF(ACSS_Clean!{A_("is_resubmission")}:{A_("is_resubmission")},TRUE)', "column is_resubmission")
line("Rows with at least one redaction", f'=COUNTIF(ACSS_Clean!{A_("redactions")}:{A_("redactions")},"?*")-1', "column redactions")
line("Dates that may have day/month swapped", f'=COUNTIF(ACSS_Clean!{A_("date_confidence")}:{A_("date_confidence")},"medium")', "date_confidence = medium")
r += 1
head("What the ACSS rows actually are")
for cls in ["portal_technical_support", "other_unclear", "policy_or_approval_enquiry", "faculty_record_dispute",
            "feedback_appreciation_suggestion", "grievance_against_institution"]:
    line(cls, f'=COUNTIF(ACSS_Clean!{A_("record_class")}:{A_("record_class")},"{cls}")', "share: see column D")
    sm.cell(r - 1, 4, f"=B{r-1}/$B${KEPT_ROW}").number_format = "0.0%"
line("Distinct genuine grievance matters", int(stats["distinct_genuine_grievance_matters"]), "after grouping resubmissions (complaint_group)")
r += 1
head("Synthetic monitoring layer")
line("Institutions", "=COUNTA(Institutions!A:A)-1")
line("Cases", "=COUNTA(Cases!A:A)-1")
for st in ["Received", "Assigned", "Under Process", "Action Initiated", "Disposed"]:
    line(f"  status: {st}", f'=COUNTIF(Cases!{C_("status")}:{C_("status")},"{st}")')
line("  overdue (active and past expected disposal)", f'=COUNTIF(Cases!{C_("is_overdue")}:{C_("is_overdue")},TRUE)')
line("  transferred at least once", f'=COUNTIF(Cases!{C_("was_transferred")}:{C_("was_transferred")},TRUE)')
line("  institution identity not yet verified", f'=COUNTA(Cases!{C_("identity_status")}:{C_("identity_status")})-1-COUNTIF(Cases!{C_("identity_status")}:{C_("identity_status")},"VERIFIED")', "excluded from overlap detection")
for b in ["PGRC", "RB", "AB"]:
    line(f"  Bureau {b}", f'=COUNTIF(Cases!{C_("bureau")}:{C_("bureau")},"{b}")')
line("Case actions", "=COUNTA(Case_Actions!A:A)-1")
line("Labelled pairs", "=COUNTA(Overlap_Ground_Truth!A:A)-1")
line("  same_matter", f'=COUNTIF(Overlap_Ground_Truth!{G_("relation")}:{G_("relation")},"same_matter")')
line("  different_matter", f'=COUNTIF(Overlap_Ground_Truth!{G_("relation")}:{G_("relation")},"different_matter")')
line("  duplicate_submission", f'=COUNTIF(Overlap_Ground_Truth!{G_("relation")}:{G_("relation")},"duplicate_submission")')
line("  pairs that should raise a cross-Bureau alert", f'=COUNTIF(Overlap_Ground_Truth!{G_("expected_cross_bureau_alert")}:{G_("expected_cross_bureau_alert")},TRUE)', "cross-Bureau, both active, ≤180 days apart, related")
sm.sheet_properties.tabColor = "1F3A5F"

# ------------------------------------------------------------------ data dictionary
DD = [
    ("ACSS_Clean", "acss_id", "Stable ID = ACSS- + original Excel row number"), ("ACSS_Clean", "source_row", "Row number in ACSSv1.xlsx"),
    ("ACSS_Clean", "created_on", "Submission timestamp (both source formats parsed)"),
    ("ACSS_Clean", "date_confidence", "high = unambiguous; medium = day ≤ 12, may be day/month swapped by Excel"),
    ("ACSS_Clean", "type / module / nature", "Original labels from the source (noisy: e.g. many 'Grievance' rows are portal help requests)"),
    ("ACSS_Clean", "subject / description", "Original text with personal and identifying details replaced by [PERSON], [INSTITUTION], [EMAIL], [PHONE], [PAN], [IFSC], [BANK_AC], [AICTE_ID], [ADDRESS], [URL], [LONG_NUMBER]"),
    ("ACSS_Clean", "desc_len / embed_eligible", "Description length; TRUE if ≥ 40 characters (long enough for semantic similarity)"),
    ("ACSS_Clean", "record_class", "What the row really is: portal_technical_support, policy_or_approval_enquiry, faculty_record_dispute, grievance_against_institution, feedback_appreciation_suggestion, other_unclear"),
    ("ACSS_Clean", "class_method / review_status", "How record_class was set; whether a person reviewed it"),
    ("ACSS_Clean", "suggested_demo_use", "How the row can honestly be used in the prototype"),
    ("ACSS_Clean", "complaint_group / is_resubmission", "Same complaint resubmitted with small edits share a group; later copies flagged TRUE"),
    ("ACSS_Clean", "redactions / had_aicte_id", "Which redactions were applied; whether the text cited an AICTE ID"),
    ("ACSS_Clean", "dup_group / is_repeat_submission / row_hash", "Exact-text grouping and hash for import de-duplication"),
    ("ACSS_Clean", "provenance / source_dataset / source_doi", "IMPORTED_UNVERIFIED + citation"),
    ("Institutions", "institution_id / aicte_permanent_id", "Fictional IDs (SYN-INST-…, SYN-1-…). Never real AICTE IDs"),
    ("Institutions", "name, stream, state, district, region, ownership, established", "Fictional institution on a real place; region grouping is synthetic"),
    ("Identity_Tests", "raw_name, raw_location, expected_institution_id, expected_decision, case_note", "Input a raw reference and check the identity engine returns the expected decision"),
    ("Bureaus", "code, name, sla_days_default, note", "Bureau master; SLA days synthetic"),
    ("Users", "user_id, display_name, role, bureau, region, email", "Role-numbered demo accounts; e-mails use the reserved .invalid domain"),
    ("Category_Matrix", "category_a, bureau_a, category_b, bureau_b, rule_points", "Related categories across Bureaus; used by the rule score"),
    ("Cases", "case_id … expected_disposal_date", "All mandatory case-tracking fields from the problem statement"),
    ("Cases", "identity_status / raw_institution_name", "VERIFIED, or VERIFICATION_QUEUE / UNRESOLVED with the raw name as received"),
    ("Cases", "is_overdue / was_transferred", "Derived flags as of 2026-10-31"),
    ("Cases", "refs", "Reference numbers cited (letter numbers, other case IDs); used by the shared-reference rule"),
    ("Cases", "complainant_ref / complainant_hash", "Fictional complainant key and its salted hash (the app should store only the hash)"),
    ("Cases", "matter_id / design_note", "Generator bookkeeping: which cases were designed as the same matter. Hide from the engine; use only for evaluation"),
    ("Case_Actions", "case_id, ts, actor, action_type, remarks, from_status, to_status", "Append-only history: created, assigned, status_change, document, transfer, escalation"),
    ("Transfers", "case_id, from_bureau, to_bureau, reason, requested_by/at, accepted_by/at, state", "Two-step Bureau transfers"),
    ("Overlap_Ground_Truth", "relation", "same_matter / different_matter / duplicate_submission"),
    ("Overlap_Ground_Truth", "expected_min_band", "HIGH if a reference is shared or it is a duplicate; MEDIUM for other same-matter pairs; NONE for different matters"),
    ("Overlap_Ground_Truth", "expected_cross_bureau_alert", "TRUE if a correct engine should raise a coordination alert (cross-Bureau, both active, ≤180 days apart, related)"),
    ("All sheets", "provenance", "VERIFIED / IMPORTED_UNVERIFIED / SYNTHETIC / DERIVED"),
]
dd = wb.create_sheet("Data_Dictionary")
dd.append(["Sheet", "Column(s)", "Meaning"])
for row in DD:
    dd.append(list(row))
for c in dd[1]:
    c.font, c.fill = HDR, HFILL
for L, w in zip("ABC", (20, 44, 100)):
    dd.column_dimensions[L].width = w
for row in dd.iter_rows(min_row=2):
    for c in row:
        c.alignment = WRAP
dd.freeze_panes = "A2"

# ------------------------------------------------------------------ data sheets
data_sheet("ACSS_Clean", acss, fill="F59E0B", wide_cols=("subject", "description", "suggested_demo_use"))
data_sheet("ACSS_Grievances", griev[["acss_id", "created_on", "type", "nature", "subject", "description", "complaint_group", "is_resubmission", "review_status", "provenance"]],
           fill="F59E0B", wide_cols=("subject", "description"))
data_sheet("Institutions", syn["institutions"], fill="7C3AED")
data_sheet("Identity_Tests", syn["institution_alias_tests"], fill="7C3AED")
data_sheet("Bureaus", syn["bureaus"], fill="7C3AED", wide_cols=("note",))
data_sheet("Users", syn["users"], fill="7C3AED")
data_sheet("Category_Matrix", syn["category_matrix"], fill="7C3AED")
data_sheet("Cases", syn["cases"], fill="7C3AED", wide_cols=("description",))
data_sheet("Case_Actions", syn["case_actions"], fill="7C3AED")
data_sheet("Transfers", syn["transfers"], fill="7C3AED")
data_sheet("Overlap_Ground_Truth", syn["overlap_ground_truth"], fill="7C3AED", wide_cols=("design_note",))
wb.save(f"{OUT}/SIH2026_AICTE_prototype_dataset.xlsx")

# ------------------------------------------------------------------ CSV bundle (for mongoimport / scripts)
csvs = {"acss_clean": acss, "acss_grievances": griev, **{k: v for k, v in syn.items()}}
for n, df in csvs.items():
    df.to_csv(f"{OUT}/csv/{n}.csv", index=False)
with zipfile.ZipFile(f"{OUT}/SIH2026_dataset_csv_and_scripts.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for n in csvs:
        z.write(f"{OUT}/csv/{n}.csv", f"csv/{n}.csv")
    for s in ["clean_acss.py", "gen_synthetic.py", "package.py", "verify.py"]:
        z.write(f"/home/claude/dataset/scripts/{s}", f"scripts/{s}")
print("saved")
